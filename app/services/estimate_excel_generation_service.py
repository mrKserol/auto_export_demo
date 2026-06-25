from __future__ import annotations

import re
import tempfile
from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE_PATH = Path("templates/smeta_template.xlsx")
PLACEHOLDER_RE = re.compile(r"{{\s*([a-zA-Z0-9_.]+)\s*}}")


class EstimateExcelTemplateNotFoundError(FileNotFoundError):
    pass


def generate_estimate_excel(estimate: dict, specification: dict | None = None) -> Path:
    if not TEMPLATE_PATH.is_file():
        raise EstimateExcelTemplateNotFoundError(
            f"Estimate Excel template not found: {TEMPLATE_PATH}"
        )

    workbook = load_workbook(TEMPLATE_PATH)
    context = build_estimate_excel_context(estimate, specification)

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value and "}}" in cell.value:
                    cell.value = render_template_cell(cell.value, context)

    output_dir = Path(tempfile.gettempdir()) / "auto_export_demo_estimates"
    output_dir.mkdir(parents=True, exist_ok=True)

    estimate_id = estimate.get("id") or "unknown"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"estimate_{estimate_id}_{timestamp}.xlsx"

    workbook.save(output_path)
    return output_path


def build_estimate_excel_context(estimate: dict, specification: dict | None = None) -> dict:
    inspect_transport_price = estimate.get("inspect_transport_price") or 0
    exchange_rate = estimate.get("exchange_rate") or 0
    inspect_transport_price_rub = _multiply_money(inspect_transport_price, exchange_rate)

    estimates_context = {
        "price_currency": estimate.get("price_currency") or "",
        "exchange_rate": format_template_number(estimate.get("exchange_rate")),
        "price_abroad": format_template_number(estimate.get("price_abroad")),
        "price_rub": format_template_number(estimate.get("price_rub")),
        "bank_commission": format_template_number(estimate.get("bank_commission")),
        "inspect_transport_price": format_template_number(inspect_transport_price),
        "inspect_transport_price_rub": format_template_number(inspect_transport_price_rub),
        "transit_declaration_price": format_template_number(
            estimate.get("transit_declaration_price") or 0
        ),
        "insurance_shipment": format_template_number(estimate.get("insurance_shipment") or 0),
        "customs_sbor": format_template_number(estimate.get("customs_sbor") or 0),
        "customs_tax": format_template_number(estimate.get("customs_tax") or 0),
        "customs_util": format_template_number(estimate.get("customs_util") or 0),
        "customs_nds": format_template_number(estimate.get("customs_nds") or 0),
        "customs_excise": format_template_number(estimate.get("customs_excise") or 0),
        "customs_total": format_template_number(estimate.get("customs_total") or 0),
        "customs_total2": format_template_number(estimate.get("customs_total2") or 0),
        "custom_clearing": format_template_number(estimate.get("custom_clearing") or 0),
        "contractor_comission": format_template_number(estimate.get("contractor_comission") or 0),
        "contractor_comission_prepayment": format_template_number(
            estimate.get("contractor_comission_prepayment") or 0
        ),
        "contractor_comission_postpayment": format_template_number(
            estimate.get("contractor_comission_postpayment") or 0
        ),
        "total_rub": format_template_number(estimate.get("total_rub")),
    }

    specification_context = {
        "brand": "",
        "model": "",
        "year": "",
        "eng_capacity": "",
        "eng_type": "",
        "color": "",
        "complectation": "",
    }

    if specification:
        specification_context.update(
            {
                "brand": specification.get("brand") or "",
                "model": specification.get("model") or "",
                "year": specification.get("year") or "",
                "eng_capacity": specification.get("eng_capacity") or "",
                "eng_type": specification.get("eng_type") or "",
                "color": specification.get("color") or "",
                "complectation": specification.get("complectation") or "",
            }
        )

    return {
        "estimates": estimates_context,
        "specification": specification_context,
    }


def render_template_cell(value: str, context: dict) -> object:
    matches = list(PLACEHOLDER_RE.finditer(value))
    if not matches:
        return value

    stripped = value.strip()
    if len(matches) == 1 and stripped == matches[0].group(0).strip():
        key = matches[0].group(1)
        rendered = get_context_value(context, key)

        numeric = try_parse_number(rendered)
        if numeric is not None:
            return numeric

        return rendered

    def replace(match: re.Match[str]) -> str:
        key = match.group(1)
        return str(get_context_value(context, key))

    return PLACEHOLDER_RE.sub(replace, value)


def get_context_value(context: dict, dotted_key: str) -> object:
    current: object = context
    for part in dotted_key.split("."):
        if not isinstance(current, dict):
            return ""
        current = current.get(part)
        if current is None:
            return ""
    return current


def try_parse_number(value: object) -> float | int | None:
    if value is None:
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        return value

    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return int(value)
        return float(value)

    text = str(value).strip().replace(" ", "").replace(",", ".")
    if not text:
        return None

    try:
        number = Decimal(text)
    except Exception:
        return None

    if number == number.to_integral_value():
        return int(number)

    return float(number)


def format_template_number(value: object) -> str:
    if value is None:
        return ""

    try:
        number = Decimal(str(value))
    except Exception:
        return str(value)

    normalized = number.normalize()
    text = format(normalized, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text


def _multiply_money(left: object, right: object) -> Decimal:
    try:
        return Decimal(str(left)) * Decimal(str(right))
    except Exception:
        return Decimal("0")
